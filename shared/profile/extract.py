"""Extract plain text from PDF or DOCX."""

from __future__ import annotations

from pathlib import Path


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in {".docx", ".doc"}:
        return _extract_docx(path)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")
    raise ValueError(f"Unsupported CV format: {suffix}. Use PDF, DOCX, TXT, or MD.")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    text = "\n\n".join(pages).strip()
    if not text:
        raise ValueError(
            f"No text extracted from {path.name}. It may be a scanned PDF; "
            "run OCR first or paste the CV as .txt."
        )
    return text


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()